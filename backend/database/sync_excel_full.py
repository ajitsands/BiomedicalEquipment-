import openpyxl
import subprocess

excel_path = 'e:/medical_equpments/who_documents/Multispeciality_Hospital_Biomedical_Equipment_Checklists.xlsx'
sql_path = 'e:/medical_equpments/backend/database/sync_excel_full.sql'

wb = openpyxl.load_workbook(excel_path, data_only=True)

def esc(val):
    if val is None:
        return 'NULL'
    val_str = str(val).strip()
    val_str = val_str.replace('\\', '\\\\').replace("'", "''")
    return f"'{val_str}'"

sql_lines = []

sql_lines.append("-- Complete Synchronization of WHO 63 Equipment Types & 2,079 Checklist Items")
sql_lines.append("-- Powered and Engineered by SanDS Lab Middle East W.L.L")
sql_lines.append("USE `biomedical_cmms`;\n")

# 1. Alter tables
sql_lines.append("""
SET @dbname = DATABASE();
SET @tablename = 'equipment_types';

-- Add asset_type_code to equipment_types
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'asset_type_code') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN asset_type_code VARCHAR(30) AFTER id'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add department_scope
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'department_scope') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN department_scope VARCHAR(150) AFTER category'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add common_complaints
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'common_complaints') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN common_complaints TEXT AFTER description'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add routine_user_inspection
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'routine_user_inspection') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN routine_user_inspection TEXT AFTER common_complaints'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add technical_inspection_scope
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'technical_inspection_scope') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN technical_inspection_scope TEXT AFTER routine_user_inspection'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add suggested_pm_schedule
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'suggested_pm_schedule') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN suggested_pm_schedule VARCHAR(150) AFTER default_pm_frequency'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add responsible_staff
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'responsible_staff') > 0,
  'SELECT 1',
  'ALTER TABLE equipment_types ADD COLUMN responsible_staff VARCHAR(150) AFTER technical_inspection_scope'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add checklist_phase to checklist_items
SET @tablename = 'checklist_items';
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'checklist_phase') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN checklist_phase VARCHAR(100) DEFAULT \"Preventive maintenance / technical inspection\" AFTER category'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add expected_result
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'expected_result') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN expected_result TEXT AFTER instructions'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add responsible_role
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'responsible_role') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN responsible_role VARCHAR(100) AFTER expected_result'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add suggested_timing
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'suggested_timing') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN suggested_timing VARCHAR(100) AFTER responsible_role'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add result_options
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'result_options') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN result_options VARCHAR(100) DEFAULT \"Pass / Fail / N/A\" AFTER suggested_timing'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;

-- Add failure_action
SET @preparedStatement = (SELECT IF(
  (SELECT COUNT(*) FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_SCHEMA = @dbname AND TABLE_NAME = @tablename AND COLUMN_NAME = 'failure_action') > 0,
  'SELECT 1',
  'ALTER TABLE checklist_items ADD COLUMN failure_action TEXT AFTER result_options'
));
PREPARE stmt FROM @preparedStatement; EXECUTE stmt; DEALLOCATE PREPARE stmt;
""")

# 2. Complaint Codes
sql_lines.append("""
CREATE TABLE IF NOT EXISTS `complaint_codes` (
    `id` INT AUTO_INCREMENT PRIMARY KEY,
    `code` VARCHAR(30) NOT NULL UNIQUE,
    `category` VARCHAR(100) NOT NULL,
    `description` TEXT NOT NULL,
    `initial_owner` VARCHAR(150),
    `priority_guidance` VARCHAR(100),
    `created_at` TIMESTAMP DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
""")

cc_sheet = wb['Complaint Codes']
sql_lines.append("-- 2. Populate Complaint Codes")
for row in cc_sheet.iter_rows(min_row=2, values_only=True):
    code, cat, desc, owner, priority = row[0], row[1], row[2], row[3], row[4]
    if code:
        sql = f"INSERT INTO complaint_codes (code, category, description, initial_owner, priority_guidance) VALUES ({esc(code)}, {esc(cat)}, {esc(desc)}, {esc(owner)}, {esc(priority)}) ON DUPLICATE KEY UPDATE category=VALUES(category), description=VALUES(description), initial_owner=VALUES(initial_owner), priority_guidance=VALUES(priority_guidance);"
        sql_lines.append(sql)

# 3. Handle Equipment Types:
# Clean older checklist items/templates to reload the pure 2,079 items from the Excel sheet
sql_lines.append("""
-- 3. Prepare clean templates and items
-- Temporarily disable foreign key checks for checklist rebuild
SET FOREIGN_KEY_CHECKS = 0;
TRUNCATE TABLE checklist_items;
DELETE FROM checklist_templates;
""")

em_sheet = wb['Equipment Master']
sql_lines.append("-- Populate 63 Equipment Master Types")

for row in em_sheet.iter_rows(min_row=2, values_only=True):
    code = row[0]
    dept_scope = row[1]
    name = row[2]
    common_faults = row[3]
    routine_insp = row[4]
    tech_pm_scope = row[5]
    maint_type = row[6]
    sched = row[7]
    resp_staff = row[8]
    risk = row[9]

    if not name:
        continue

    # Map general category based on department or name
    cat = 'General Clinical Medical'
    dept_lower = (dept_scope or '').lower()
    name_lower = name.lower()

    if any(k in dept_lower or k in name_lower for k in ['icu', 'ventilator', 'defibrillator', 'resuscitation']):
        cat = 'Life Support'
    elif any(k in dept_lower or k in name_lower for k in ['theatre', 'surgical', 'anaesthesia', 'diathermy', 'c-arm', 'tourniquet']):
        cat = 'Surgical & Anaesthesia'
    elif any(k in dept_lower or k in name_lower for k in ['radiology', 'x-ray', 'ct', 'mri', 'ultrasound', 'mammography']):
        cat = 'Medical Imaging & Radiology'
    elif any(k in dept_lower or k in name_lower for k in ['lab', 'analyser', 'microscope', 'centrifuge']):
        cat = 'Laboratory & Diagnostics'
    elif any(k in dept_lower or k in name_lower for k in ['cssd', 'steriliser', 'autoclave', 'disinfector']):
        cat = 'Sterilization & Decontamination'
    elif any(k in dept_lower or k in name_lower for k in ['dialysis', 'haemodialysis']):
        cat = 'Renal Care & Dialysis'
    elif any(k in dept_lower or k in name_lower for k in ['nicu', 'infant', 'incubator', 'warmer', 'fetal', 'maternity']):
        cat = 'Maternal & Neonatal Care'
    elif any(k in dept_lower or k in name_lower for k in ['workshop', 'simulator']):
        cat = 'Biomedical Workshop & Testing'
    elif any(k in dept_lower or k in name_lower for k in ['cardiology', 'ecg', 'stress test', 'holter']):
        cat = 'Cardiology'
    elif any(k in dept_lower or k in name_lower for k in ['dental']):
        cat = 'Dental Care'
    elif any(k in dept_lower or k in name_lower for k in ['physiotherapy', 'tens']):
        cat = 'Physiotherapy & Rehab'
    elif any(k in dept_lower or k in name_lower for k in ['pharmacy', 'refrigerator', 'blood bank']):
        cat = 'Cold Chain & Storage'

    pm_freq = 'Quarterly'
    if sched:
        s_lower = sched.lower()
        if '12 month' in s_lower or 'annual' in s_lower:
            pm_freq = 'Annual'
        elif '6 month' in s_lower:
            pm_freq = 'Semi-Annual'
        elif 'monthly' in s_lower:
            pm_freq = 'Monthly'

    risk_formatted = risk or 'High'
    if 'Critical' in risk_formatted:
        risk_class = 'Critical (Class III)'
    elif 'High' in risk_formatted:
        risk_class = 'High (Class IIb)'
    elif 'Medium' in risk_formatted:
        risk_class = 'Medium (Class IIa)'
    else:
        risk_class = 'Low (Class I)'

    # Upsert by name
    sql = f"""
INSERT INTO equipment_types (
    asset_type_code, name, category, department_scope, risk_classification,
    default_pm_frequency, suggested_pm_schedule, standard_regulatory_code,
    description, common_complaints, routine_user_inspection, technical_inspection_scope, responsible_staff
) VALUES (
    {esc(code)}, {esc(name)}, {esc(cat)}, {esc(dept_scope)}, {esc(risk_class)},
    {esc(pm_freq)}, {esc(sched)}, 'WHO Technical Series / IEC 62353',
    {esc(maint_type)}, {esc(common_faults)}, {esc(routine_insp)}, {esc(tech_pm_scope)}, {esc(resp_staff)}
) ON DUPLICATE KEY UPDATE
    asset_type_code = VALUES(asset_type_code),
    category = VALUES(category),
    department_scope = VALUES(department_scope),
    risk_classification = VALUES(risk_classification),
    default_pm_frequency = VALUES(default_pm_frequency),
    suggested_pm_schedule = VALUES(suggested_pm_schedule),
    description = VALUES(description),
    common_complaints = VALUES(common_complaints),
    routine_user_inspection = VALUES(routine_user_inspection),
    technical_inspection_scope = VALUES(technical_inspection_scope),
    responsible_staff = VALUES(responsible_staff);
"""
    sql_lines.append(sql.strip())

# Remove any old equipment_types that are not among the 63 in Excel (if any duplicates)
sql_lines.append("""
-- Keep only the official 63 equipment types
DELETE FROM equipment_types WHERE asset_type_code IS NULL;
""")

# 4. Insert 63 Templates (one per equipment type)
sql_lines.append("""
-- 4. Insert 63 Checklist Templates
INSERT INTO checklist_templates (equipment_type_id, title, frequency, standard_reference, estimated_minutes, description)
SELECT 
    id, 
    CONCAT(name, ' Inspection & PM Protocol (33-Point WHO Checklist)'), 
    default_pm_frequency, 
    'WHO Medical Device Technical Series & IEC 62353', 
    45, 
    CONCAT('Comprehensive 33-point verification covering User Pre-Use Checks (1-10), Preventive Maintenance Technical Inspection (1-13), and Corrective Maintenance Breakdown Testing (1-10) for ', name)
FROM equipment_types;
""")

# 5. Insert all 2,079 Checklist Items
cl_sheet = wb['Checklist Library']
sql_lines.append("-- 5. Insert all 2,079 Checklist Items")

for row in cl_sheet.iter_rows(min_row=2, values_only=True):
    eq_type = row[0]
    phase = row[1]
    item_no = row[2]
    task = row[3]
    expected = row[4]
    role = row[5]
    timing = row[6]
    options = row[7]
    failure_action = row[8]

    if not eq_type or not task:
        continue

    # Determine category
    task_lower = task.lower()
    if 'ground' in task_lower or 'leakage' in task_lower or 'earth' in task_lower or 'electrical safety' in task_lower:
        cat = 'Electrical Safety Test'
        check_type = 'numeric'
    elif 'calibration' in task_lower or 'accuracy' in task_lower or 'tolerance' in task_lower or 'simulator' in task_lower or 'pressure' in task_lower:
        cat = 'Functional & Calibration Verification'
        check_type = 'pass_fail'
    elif 'alarm' in task_lower or 'warning' in task_lower:
        cat = 'Alarms & Warning System'
        check_type = 'pass_fail'
    elif 'battery' in task_lower or 'power' in task_lower or 'mains' in task_lower:
        cat = 'Power & Battery Backup'
        check_type = 'pass_fail'
    elif 'damage' in task_lower or 'cable' in task_lower or 'serial' in task_lower or 'visual' in task_lower or 'casing' in task_lower or 'clean' in task_lower:
        cat = 'Physical & Visual Inspection'
        check_type = 'pass_fail'
    else:
        cat = 'Technical Performance Inspection'
        check_type = 'pass_fail'

    sql = f"""
INSERT INTO checklist_items (
    template_id, category, checklist_phase, step_number, title, instructions,
    expected_result, responsible_role, suggested_timing, result_options, failure_action,
    check_type, is_mandatory
)
SELECT 
    ct.id,
    {esc(cat)},
    {esc(phase)},
    {int(item_no or 1)},
    {esc(task)},
    {esc(expected)},
    {esc(expected)},
    {esc(role)},
    {esc(timing)},
    {esc(options or 'Pass / Fail / N/A')},
    {esc(failure_action)},
    {esc(check_type)},
    1
FROM checklist_templates ct
JOIN equipment_types et ON ct.equipment_type_id = et.id
WHERE et.name = {esc(eq_type)};
"""
    sql_lines.append(sql.strip())

# 6. Re-map seeded equipment to matching new equipment types
sql_lines.append("""
-- 6. Re-map existing equipments to matching equipment type IDs
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Ventilator, ICU' LIMIT 1) WHERE asset_id LIKE '%VENT%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Defibrillator / Monitor' LIMIT 1) WHERE asset_id LIKE '%DEF%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Multiparameter Patient Monitor' LIMIT 1) WHERE asset_id LIKE '%MON%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Infusion Pump' LIMIT 1) WHERE asset_id LIKE '%INFP%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Syringe Pump' LIMIT 1) WHERE asset_id LIKE '%SYRP%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Haemodialysis Machine' LIMIT 1) WHERE asset_id LIKE '%DIAL%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Electrosurgical Unit / Diathermy' LIMIT 1) WHERE asset_id LIKE '%ESU%';
UPDATE equipments SET equipment_type_id = (SELECT id FROM equipment_types WHERE name = 'Anaesthesia Workstation' LIMIT 1) WHERE asset_id LIKE '%ANES%';

-- Re-enable foreign key checks
SET FOREIGN_KEY_CHECKS = 1;
""")

with open(sql_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(sql_lines))

print(f"Generated {sql_path} with {len(sql_lines)} SQL statements.")
